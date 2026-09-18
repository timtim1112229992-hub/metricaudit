"""P9: the tables, rendered from the results file.

Same discipline as the figures. A table built from the results cannot show a
number the results do not contain, and the alternative, building tables from the
frames, is how a paper ends up with a table and a sentence that disagree.
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd

from .rules import RULES


def _fmt(value, places: int = 3) -> str:
    if value is None:
        return "\u2013"
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float):
        if pd.isna(value):
            return "\u2013"
        if value == int(value) and abs(value) < 1e9:
            return f"{int(value)}"
        return f"{value:.{places}f}"
    return str(value)


def _frame(results: dict, key: str) -> pd.DataFrame:
    return pd.DataFrame(results.get(key) or [])


def table_counting_rules() -> pd.DataFrame:
    """The declared register, as it was committed."""
    return pd.DataFrame([{
        "Reported column": r.column,
        "Kind": r.kind,
        "Source": r.source,
        "Counting rule": r.statement,
        "Deduplication assumption": r.dedup,
        "Inclusion window": r.window,
        "Records lacking an identifier": r.missing,
    } for r in RULES])


def table_census(results: dict) -> pd.DataFrame:
    census = results["census"]
    parity = _frame(results, "parity_identifiers")
    composite = _frame(results, "parity_composite")

    rows = []
    for _, row in parity.iterrows():
        match = composite[composite["event_category"] == row["event_category"]]
        rows.append({
            "Event category": row["event_category"],
            "Operational store": int(row["n_operational"]),
            "Reporting store": int(row["n_reporting"]),
            "Difference": int(row["count_difference"]),
            "Shared identifiers": int(row["n_shared_identifiers"]),
            "Jaccard on identifiers": _fmt(float(row["jaccard"])),
            "Pairs matched on a composite key": (int(match["n_matched_pairs"].iloc[0])
                                                 if len(match) else "\u2013"),
            "Unmatched": (int(match["operational_unmatched"].iloc[0])
                          if len(match) else "\u2013"),
        })
    for category, count in (census["exclusive_to_reporting"] or {}).items():
        rows.append({"Event category": f"{category} (reporting store only)",
                     "Operational store": 0, "Reporting store": int(count),
                     "Difference": int(count), "Shared identifiers": 0,
                     "Jaccard on identifiers": "\u2013",
                     "Pairs matched on a composite key": "\u2013",
                     "Unmatched": "\u2013"})
    rows.append({"Event category": "All categories",
                 "Operational store": int(census["n_operational"]),
                 "Reporting store": int(census["n_reporting"]),
                 "Difference": int(census["raw_difference"]),
                 "Shared identifiers": 0, "Jaccard on identifiers": "\u2013",
                 "Pairs matched on a composite key": "\u2013",
                 "Unmatched": "\u2013"})
    return pd.DataFrame(rows)


def table_reconciliation(results: dict) -> pd.DataFrame:
    frame = _frame(results, "reconciliation_by_indicator")
    return pd.DataFrame([{
        "Reported column": row["indicator"],
        "Kind": row["kind"],
        "Groups": int(row["n_groups"]),
        "Reconciling": int(row["n_matching"]),
        "Diverging": int(row["n_diverging"]),
        "Outcome": row["classification"],
        "Note": row["reason"] or "",
    } for _, row in frame.iterrows()])


def table_divergence_by_group(results: dict) -> pd.DataFrame:
    frame = _frame(results, "reconciliation_by_group")
    if frame.empty:
        return frame
    frame = frame[frame["numeric"].astype(bool)]
    divergent = set(_frame(results, "reconciliation_by_indicator")
                    .query("classification == 'divergent'")["indicator"])
    frame = frame[frame["indicator"].isin(divergent)].sort_values(
        ["indicator", "group_number"])
    return pd.DataFrame([{
        "Reported column": row["indicator"],
        "Group": int(row["group_number"]),
        "Reported": _fmt(row["reported"]),
        "Recomputed": _fmt(row["recomputed"]),
        "Difference": _fmt(row["absolute_difference"]),
        "Ratio": _fmt(row["ratio"], 2),
        "Direction": row["direction"],
    } for _, row in frame.iterrows()])


def table_attribution(results: dict) -> pd.DataFrame:
    """Attribution for every divergent indicator and group.

    Residuals for a count are counts and residuals for a timestamp are seconds,
    so the unit travels in the cell. A single unitless residual column would
    read as though a three second offset and a difference of three clicks were
    the same size of miss.
    """
    frame = _frame(results, "attribution")
    if frame.empty:
        return frame
    kinds = {r.column: r.kind for r in RULES}
    rows = []
    for _, row in frame.iterrows():
        temporal = kinds.get(row["indicator"]) == "derived" and pd.isna(row["reported"])
        residual = row["residual"]
        rows.append({
            "Reported column": row["indicator"],
            "Group": int(row["group_number"]),
            "Reported": _fmt(row["reported"]),
            "Mechanism": row["mechanism"],
            "Predicted": _fmt(row["predicted"]),
            "Residual": ("\u2013" if residual is None or pd.isna(residual)
                         else f"{float(residual):.3f} s" if temporal
                         else _fmt(float(residual), 3)),
            "Note": row["reason"] or "",
        })
    return pd.DataFrame(rows)


def table_predictions(results: dict) -> pd.DataFrame:
    frame = _frame(results, "prediction_matrix")
    if frame.empty:
        return frame
    ordered = ["group_number", "reported", "single_store", "cross_store_summation",
               "repeated_client_emission", "absent_idempotency", "gauge_summation",
               "gauge_final_value", "decision_count"]
    present = [c for c in ordered if c in frame.columns]
    out = frame[present].sort_values("group_number")
    out.columns = [c.replace("_", " ").replace("group number", "Group")
                   .capitalize() if c != "group_number" else "Group"
                   for c in out.columns]
    return out


def table_prediction_totals(results: dict) -> pd.DataFrame:
    """What each candidate predicts in total, beside the reported figure.

    The per-group matrix carries the evidence and is too wide to read. This
    collapses it to one row per candidate, which is the form in which the
    separation between candidates is obvious: either a candidate lands on the
    reported total or it is nowhere near it.
    """
    frame = _frame(results, "prediction_matrix")
    if frame.empty:
        return frame
    candidates = [c for c in frame.columns
                  if c not in ("indicator", "group_id", "group_number", "reported")]
    reported = float(frame["reported"].sum())

    rows = [{
        "Source of the figure": "Reported by the view",
        "Total across groups": _fmt(reported),
        # Through the same formatter as every other ratio in the column. Writing
        # the value out here instead gives the reference row a different number
        # of decimals from the candidate that matches it exactly, which reads as
        # a difference between them.
        "Ratio to the reported total": _fmt(1.0),
    }]
    for name in sorted(candidates):
        total = float(frame[name].sum())
        rows.append({
            "Source of the figure": name.replace("_", " "),
            "Total across groups": _fmt(total),
            "Ratio to the reported total": _fmt(total / reported) if reported else "",
        })
    return pd.DataFrame(rows)


def table_gauge_profile(results: dict) -> pd.DataFrame:
    """The shape of the field the attributed mechanism sums.

    Carried as a table because the mechanism's name implies a claim about the
    field, and a reader should be able to check that claim rather than take it.
    """
    frame = _frame(results, "gauge_profile")
    if frame.empty:
        return frame
    return pd.DataFrame([{
        "Group": int(row["group_number"]),
        "Snapshots carrying the counter": int(row["n_snapshots"]),
        "Runs": int(row["n_runs"]),
        "Resets": int(row["n_resets"]),
        "Resets at a change of stage": int(row["resets_at_a_stage_change"]),
        "Longest run": int(row["longest_run"]),
        "Highest value reached": _fmt(row["highest_value"]),
        "Value at the last snapshot": _fmt(row["final_value"]),
        "Summed across snapshots": _fmt(row["summed"]),
        "Accumulates across the session": _fmt(bool(row["accumulates_across_the_session"])),
    } for _, row in frame.iterrows()])


def table_temporal_trace(results: dict) -> pd.DataFrame:
    """Which store each group's reported activity time was stamped from, and
    what kind of record it was.

    One row per group, with the offset to each candidate store's latest record.
    A negative offset means the store holds something later than the column
    claims, which is what identifies the store the column was not stamped from.
    """
    frame = _frame(results, "temporal_trace")
    if frame.empty:
        return frame
    rows = []
    for group, block in frame.groupby("group_number"):
        row = {"Group": int(group)}
        for _, entry in block.iterrows():
            store = str(entry["store"]).replace("decisions", "agent decisions")
            row[f"Offset to latest {store} record, s"] = _fmt(entry["offset_s"], 3)
            row[f"Category of that {store} record"] = entry["latest_record_category"]
        rows.append(row)
    return pd.DataFrame(rows)


def table_completeness(results: dict) -> pd.DataFrame:
    frame = _frame(results, "completeness")
    return pd.DataFrame([{
        "Store": row["store"],
        "Check": row["check"],
        "Affected": int(row["n_affected"]),
        "Records": int(row["n_records"]),
        "Rate": _fmt(float(row["rate"]), 4),
        "95% interval": f"{float(row['ci_lower']):.4f} to {float(row['ci_upper']):.4f}",
        "One-sided bound where zero": (_fmt(float(row["zero_bound"]), 4)
                                       if row.get("zero_bound") is not None
                                       else "\u2013"),
    } for _, row in frame.iterrows()])


def table_sweep(results: dict) -> pd.DataFrame:
    frame = _frame(results, "sweep")
    if frame.empty:
        return frame
    return pd.DataFrame([{
        "Assumption": row["assumption"],
        "Setting": row["setting"],
        "Reported column": row["indicator"],
        "Groups reconciling": f"{int(row['n_matching'])} of {int(row['n_groups'])}",
        "Outcome": row["classification"],
        "Total recomputed": _fmt(row["total_recomputed"]),
    } for _, row in frame.iterrows()])


def table_dependence(results: dict) -> pd.DataFrame:
    frame = _frame(results, "assumption_dependence")
    if frame.empty:
        return frame
    return pd.DataFrame([{
        "Reported column": row["indicator"],
        "Settings examined": int(row["n_settings"]),
        "Outcomes seen": row["classifications_seen"],
        "Assumption dependent": _fmt(bool(row["assumption_dependent"])),
        "Settings under which it reconciles": row["settings_reconciling"],
    } for _, row in frame.iterrows()])


def table_consequence(results: dict) -> pd.DataFrame:
    consequence = results.get("consequence") or {}
    rows = []
    for indicator, block in sorted(consequence.items()):
        imp = block["impact"]
        for label, key in (("All groups, reported", "reported"),
                           ("All groups, reconciled", "reconciled"),
                           (f"Excluding group {imp['extreme_group']}, reported",
                            "reported_without_extreme"),
                           (f"Excluding group {imp['extreme_group']}, reconciled",
                            "reconciled_without_extreme")):
            d = imp[key]
            rows.append({"Reported column": indicator, "Value set": label,
                         "n": int(d["n"]), "Mean": _fmt(d["mean"], 2),
                         "SD": _fmt(d["sd"], 2), "Median": _fmt(d["median"], 2),
                         "Minimum": _fmt(d["min"], 1), "Maximum": _fmt(d["max"], 1)})
    return pd.DataFrame(rows)


def table_divergence_estimate(results: dict) -> pd.DataFrame:
    """M4 and M5 for each divergent indicator.

    Twelve groups is too few for an interval that assumes a shape, so the
    interval is the accelerated bootstrap and the row says which method produced
    it. The rank agreement is here because a high value is the warning: the
    ordering of groups mostly survives, so an analyst checking only who ranks
    where would find nothing wrong with values that are wrong by a factor.
    """
    rows = []
    for indicator, block in sorted((results.get("consequence") or {}).items()):
        interval = block.get("divergence_interval") or {}
        agree = block.get("agreement") or {}
        impact = block.get("impact") or {}
        mean_ratio = interval.get("mean_ratio") or {}
        median_ratio = interval.get("median_ratio") or {}
        rank = impact.get("rank_agreement") or {}
        rows.append({
            "Reported column": indicator,
            "Groups": int(interval.get("n_groups", 0)),
            "Mean ratio": _fmt(mean_ratio.get("estimate"), 2),
            "95% interval on the mean ratio":
                f"{mean_ratio['lower']:.2f} to {mean_ratio['upper']:.2f}"
                if mean_ratio.get("lower") is not None else "\u2013",
            "Median ratio": _fmt(median_ratio.get("estimate"), 2),
            "95% interval on the median ratio":
                f"{median_ratio['lower']:.2f} to {median_ratio['upper']:.2f}"
                if median_ratio.get("lower") is not None else "\u2013",
            "Interval method": mean_ratio.get("method", "\u2013"),
            "Mean agreement, as a factor":
                _fmt(10 ** agree["bias"], 2) if agree.get("estimable") else "\u2013",
            "Limits of agreement, as factors":
                f"{10 ** agree['lower']:.2f} to {10 ** agree['upper']:.2f}"
                if agree.get("estimable") else "\u2013",
            "Rank agreement, Kendall tau": _fmt(rank.get("tau"), 3),
        })
    return pd.DataFrame(rows)


def table_association(results: dict) -> pd.DataFrame:
    """M10, kept apart from the descriptive table.

    A rank correlation does not belong in a column headed mean, and putting it
    there invites a reader to compare it with one.
    """
    rows = []
    for indicator, block in sorted((results.get("consequence") or {}).items()):
        assoc = block.get("association") or {}
        if not assoc:
            continue
        for label, key in (("Reported values", "from_reported"),
                           ("Reconciled values", "from_reconciled")):
            d = assoc[key]
            rows.append({
                "Reported column": indicator,
                "Outcome": assoc["outcome"],
                "Values used": label,
                "Groups": int(d["n"]),
                "Spearman rho": _fmt(float(d["rho"]), 3),
                "p": _fmt(float(d["p"]), 3),
            })
        rows.append({
            "Reported column": indicator, "Outcome": assoc["outcome"],
            "Values used": "Change in the estimate",
            "Groups": int(assoc["from_reported"]["n"]),
            "Spearman rho": _fmt(float(assoc["from_reconciled"]["rho"])
                                 - float(assoc["from_reported"]["rho"]), 3),
            "p": "\u2013"})
    return pd.DataFrame(rows)


def table_archive(results: dict) -> pd.DataFrame:
    collections = results.get("archive_collections") or {}
    redundancy = _frame(results, "archive_redundancy")
    rows = [{"Collection": name, "Records": int(count),
             "Subset of": "no other collection",
             "Proportion of the superset it covers": "\u2013"}
            for name, count in sorted(collections.items())]
    for _, row in redundancy.iterrows():
        for entry in rows:
            if entry["Collection"] == row["subset"]:
                entry["Subset of"] = row["superset"]
                entry["Proportion of the superset it covers"] = \
                    _fmt(float(row["proportion_covered"]), 4)
    return pd.DataFrame(rows)


TABLES = (
    ("Table 1. Declared counting rules and their assumptions",
     lambda r: table_counting_rules()),
    ("Table 2. Record counts by store and event category, with parity outcome",
     table_census),
    ("Table 3. Reconciliation outcome for every column of the reporting view",
     table_reconciliation),
    ("Table 4. Divergence by group for each divergent column",
     table_divergence_by_group),
    ("Table 5. Divergence estimates and agreement, with bootstrap intervals",
     table_divergence_estimate),
    ("Table 6. Candidate mechanism predictions per group", table_predictions),
    ("Table 7. Mechanism attribution", table_attribution),
    ("Table 8. What each candidate mechanism predicts in total",
     table_prediction_totals),
    ("Table 9. Shape of the counter the attributed mechanism sums",
     table_gauge_profile),
    ("Table 10. Store and record category behind each reported activity time",
     table_temporal_trace),
    ("Table 11. Completeness and referential integrity", table_completeness),
    ("Table 12. Reconciliation under alternative assumptions", table_sweep),
    ("Table 13. Assumption dependence by column", table_dependence),
    ("Table 14. Group-level statistics under reported and reconciled values",
     table_consequence),
    ("Table 15. Association with an outcome, under reported and reconciled values",
     table_association),
    ("Table 16. Collections in the archived export and their subset relations",
     table_archive),
)


def render_docx(results_path: Path, out_path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    results = json.loads(results_path.read_text(encoding="utf-8"))
    document = Document()
    document.add_heading("Tables", level=1)

    for caption, builder in TABLES:
        frame = builder(results)
        if frame is None or frame.empty:
            continue
        document.add_heading(caption, level=2)
        table = document.add_table(rows=1, cols=len(frame.columns))
        table.style = "Table Grid"
        for cell, name in zip(table.rows[0].cells, frame.columns):
            cell.text = str(name)
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(8)
        for _, row in frame.iterrows():
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = _fmt(value)
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(8)
        document.add_paragraph()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(out_path)
    print(f"  {out_path.name}")


def render_csv(results_path: Path, directory: Path) -> None:
    results = json.loads(results_path.read_text(encoding="utf-8"))
    directory.mkdir(parents=True, exist_ok=True)
    for i, (caption, builder) in enumerate(TABLES, start=1):
        frame = builder(results)
        if frame is None or frame.empty:
            continue
        stem = f"table_{i:02d}_" + caption.split(". ", 1)[1][:44] \
            .lower().replace(" ", "_").replace(",", "")
        frame.to_csv(directory / f"{stem}.csv", index=False, encoding="utf-8")
        print(f"  {stem}.csv")
